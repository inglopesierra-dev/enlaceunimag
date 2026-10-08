"""Arma la versión del prototipo con los estudiantes REALES de BASE_DE_DATOS.xlsx.

Uso:  python construir_prototipo_real.py BASE_DE_DATOS.xlsx SALIDA.html

Toma prototipo/index.html como plantilla y le incrusta todos los estudiantes de las seis hojas de facultad, con
las 24 columnas de la base más la facultad (el nombre de la hoja), y la matrícula financiera por programa.
La página detecta esos datos y los usa en lugar de los ficticios.

El archivo de salida tiene datos personales reales, incluidos datos sensibles y de menores de edad:
- se abre localmente en el navegador y se guarda solo en el OneDrive institucional;
- no se publica como página web ni se sube a este repositorio (el .gitignore bloquea *_datos_reales.html);
- no se le agregan atenciones ni seguimientos inventados: lo que no está en la base aparece vacío.
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

AQUI = Path(__file__).resolve().parent
FACULTADES = ["INGENIERÍA", "BÁSICAS", "EMPRESARIALES Y ECONÓMICAS", "SALUD", "HUMANIDADES", "EDUCACIÓN"]
HOJA_FINANCIERA = "MATRICULA FINANCIERA 2026-2 "
FECHA_CORTE = "2026-10-07"
DIVISOR_PROMEDIO = 100   # Igual que «Cómo está mi facultad» (Configuracion!C7): 318 / 100 = 3,18
# Encabezados de la base, en orden (columnas B a Y). La columna NOMBRE(S) cambia de nombre entre hojas.
COLUMNAS_BASE = ["PERIODO", "PROGRAMA", "MODALIDAD", "CÓDIGO", "SEXO", "NOMBRES", "APELLIDOS", "MODALIDAD DE INGRESO",
                 "MATRICULADO", "CANCELACIÓN DE SEMESTRE", "CUPO ESPECIAL APLICADO", "PLAN DE ESTUDIO", "PROMEDIO ACUMULADO",
                 "TIPO DE DOCUMENTO", "NÚMERO DE DOCUMENTO", "NACIMIENTO", "ESTRATO", "ORIGEN", "DEPARTAMENTO ORIGEN",
                 "MUNICIPIO ORIGEN", "COLEGIO DE DONDE PROVIENE", "TIPO DE COLEGIO", "DEPARTAMENTO COLEGIO", "MUNICIPIO COLEGIO"]
# Cupos especiales que no revelan un dato sensible (igual que microsoft365/modelo.json).
CUPO_VISIBLE_SI_CONTIENE = ["N/A", "TALENTO", "ARTISTA", "DEPORTISTA", "BECA", "ZONA URBANA"]
# Periodo de reporte, como la tabla TPeriodos de «Cómo está mi facultad».
PERIODO_REPORTE = {"2026II": "2026-II", "20263C": "2026-II"}


def limpio(v):
    return re.sub(r"\s+", " ", v).strip() if isinstance(v, str) else v


def entero(v):
    return int(v) if isinstance(v, float) and v.is_integer() else v


def texto(v):
    v = entero(limpio(v))
    return "" if v is None else str(v)


def fecha(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=int(v))).date().isoformat()
    return ""


class Diccionario:
    """Guarda cada texto repetido una sola vez; las filas guardan su posición."""

    def __init__(self):
        self.valores, self.pos = [], {}

    def __call__(self, v):
        if v not in self.pos:
            self.pos[v] = len(self.valores)
            self.valores.append(v)
        return self.pos[v]


def leer(ruta_base):
    wb = openpyxl.load_workbook(ruta_base, read_only=True, data_only=True)
    faltan = [f for f in FACULTADES if f not in wb.sheetnames]
    if faltan:
        sys.exit(f"La base no tiene las hojas {faltan}")
    d = {k: Diccionario() for k in ("periodo", "modalidad", "sexo", "tipoDoc", "modIngreso", "origen", "lugar",
                                     "colegio", "tipoColegio", "plan")}
    programas, cupos, filas, codigos = {}, {}, [], set()
    for f_idx, fac in enumerate(FACULTADES):
        primera = True
        for r in wb[fac].iter_rows(min_col=2, max_col=25, values_only=True):
            if primera:
                encabezados = [str(x or "").strip().upper() for x in r]
                encabezados[5] = "NOMBRES"
                if [re.sub(r"\s+", " ", e) for e in encabezados] != COLUMNAS_BASE:
                    sys.exit(f"Los encabezados de la hoja {fac} no son los esperados: {encabezados}")
                primera = False
                continue
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in r):
                continue
            (periodo, programa, modalidad, codigo, sexo, nombres, apellidos, mod_ing, matric, cancel, cupo, plan, prom,
             tipo_doc, num_doc, nac, estrato, origen, dep_o, mun_o, colegio, tipo_col, dep_c, mun_c) = r
            codigo = texto(codigo)
            if codigo in codigos:
                sys.exit(f"Código repetido en la base: revisar antes de construir (hoja {fac})")
            codigos.add(codigo)
            programa = texto(programa)
            if programa not in programas:
                programas[programa] = (len(programas), f_idx)
            cupo = texto(cupo) or "N/A"
            if cupo not in cupos:
                cupos[cupo] = (len(cupos), 0 if any(k in cupo.upper() for k in CUPO_VISIBLE_SI_CONTIENE) else 1)
            prom = entero(prom)
            filas.append([
                codigo, texto(nombres), texto(apellidos), d["tipoDoc"](texto(tipo_doc)), texto(num_doc),
                d["sexo"](texto(sexo)), fecha(nac), f_idx, programas[programa][0], d["periodo"](texto(periodo)),
                d["modalidad"](texto(modalidad)), d["modIngreso"](texto(mod_ing)), texto(matric), texto(cancel),
                cupos[cupo][0], d["plan"](texto(plan)), prom if isinstance(prom, (int, float)) else None, texto(estrato),
                d["origen"](texto(origen)), d["lugar"](texto(dep_o)), d["lugar"](texto(mun_o)), d["colegio"](texto(colegio)),
                d["tipoColegio"](texto(tipo_col)), d["lugar"](texto(dep_c)), d["lugar"](texto(mun_c)),
            ])
    financiera = []
    if HOJA_FINANCIERA in wb.sheetnames:
        for prog, potencial, matric in wb[HOJA_FINANCIERA].iter_rows(min_row=3, min_col=2, max_col=4, values_only=True):
            if prog and str(prog).strip().upper() != "TOTALES" and isinstance(matric, (int, float)):
                financiera.append([texto(prog), entero(potencial) if isinstance(potencial, (int, float)) else None, int(matric)])
    wb.close()
    periodos = d["periodo"].valores
    return {
        "version": 1,
        "fuente": Path(ruta_base).name,
        "corte": FECHA_CORTE,
        "generado": dt.date.today().isoformat(),
        "divisorPromedio": DIVISOR_PROMEDIO,
        "columnasBase": ["FACULTAD (hoja)"] + COLUMNAS_BASE,
        "dic": {
            "facultad": FACULTADES,
            "programa": [[p, fi] for p, (i, fi) in sorted(programas.items(), key=lambda kv: kv[1][0])],
            "periodo": periodos,
            "periodoReporte": [PERIODO_REPORTE.get(p, p) for p in periodos],
            "modalidad": d["modalidad"].valores,
            "sexo": d["sexo"].valores,
            "tipoDoc": d["tipoDoc"].valores,
            "modIngreso": d["modIngreso"].valores,
            "cupo": [[c, s] for c, (i, s) in sorted(cupos.items(), key=lambda kv: kv[1][0])],
            "plan": d["plan"].valores,
            "origen": d["origen"].valores,
            "lugar": d["lugar"].valores,
            "colegio": d["colegio"].valores,
            "tipoColegio": d["tipoColegio"].valores,
        },
        "campos": ["codigo", "nombres", "apellidos", "tipoDoc", "documento", "sexo", "nacimiento", "facultad", "programa",
                   "periodo", "modalidad", "modIngreso", "matriculado", "cancelacion", "cupo", "plan", "promedio", "estrato",
                   "origen", "deptoOrigen", "munOrigen", "colegio", "tipoColegio", "deptoColegio", "munColegio"],
        "filas": filas,
        "matriculaFinanciera": financiera,
    }


def construir(ruta_base, salida):
    datos = leer(ruta_base)
    plantilla = (AQUI / "index.html").read_text(encoding="utf-8")
    marca = "<script>"
    if plantilla.count(marca) != 1:
        sys.exit("La plantilla debe tener un solo <script> principal")
    carga = json.dumps(datos, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = plantilla.replace(marca, f'<script id="datos-reales" type="application/json">{carga}</script>\n{marca}', 1)
    html = html.replace("<title>Bitácora de Acompañamiento</title>", "<title>Bitácora de Acompañamiento · datos reales</title>", 1)
    # El archivo se abre directo en el navegador, sin la envoltura de la plataforma: se declara HTML5 para que no use el modo de compatibilidad.
    html = '<!doctype html>\n<html lang="es">\n' + html + '\n</html>\n'
    Path(salida).write_text(html, encoding="utf-8")
    sensibles = sum(1 for f in datos["filas"] if datos["dic"]["cupo"][f[14]][1])
    print(f"{len(datos['filas'])} estudiantes, {len(datos['dic']['programa'])} programas, {sensibles} con cupo especial sensible; "
          f"{len(html) / 1e6:.1f} MB en {salida}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    construir(sys.argv[1], sys.argv[2])
